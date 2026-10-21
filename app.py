import io
import math
import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Cronoanálise & PCP", layout="wide")


def limpar(valor):
    if pd.isna(valor):
        return 0.0
    if "-" in str(valor):
        try:
            return float(pd.to_datetime(valor).day)
        except Exception:
            pass
    val_str = str(valor).replace(",", ".").strip()
    try:
        return float(val_str)
    except ValueError:
        return 0.0


st.title("⏱️ Sistema de Cronoanálise, Capacidade & OEE")
st.write(
    "Carregue a planilha da fábrica para calcular os tempos padrão e capacidades."
)

# --- CONFIGURAÇÕES NA BARRA LATERAL ---
st.sidebar.header("Parâmetros do Processo")
tolerancia_pct = st.sidebar.slider(
    "Tolerância (%)", min_value=0, max_value=50, value=15
)
tolerancia = tolerancia_pct / 100.0

jornada = st.sidebar.number_input("Jornada Diária (Horas)", value=8.0, step=0.5)

oee_pct = st.sidebar.slider(
    "OEE Estimado (%)", min_value=1, max_value=100, value=85
)
oee_estimado = oee_pct / 100.0

# --- CAIXAS DE SELEÇÃO PARA DEFINIR O QUE CALCULAR ---
st.sidebar.header("🎯 Opções de Dimensionamento")
calcular_op = st.sidebar.checkbox("Calcular Operadores", value=True)
calcular_maq = st.sidebar.checkbox("Calcular Máquinas", value=True)

meta_producao = 1000
if calcular_op or calcular_maq:
    meta_producao = st.sidebar.number_input(
        "Meta de Produção (Peças)", value=1000, step=100
    )

# Campo para o usuário digitar o nome do projeto/responsável diretamente no app
st.sidebar.header("📝 Identificação do Relatório")
nome_projeto = st.sidebar.text_input("Nome do Projeto / Responsável", value="Linha Principal")

arquivo_carregado = st.file_uploader(
    "Selecione a planilha Excel (.xlsx)", type=["xlsx"]
)

if arquivo_carregado is not None:
    df = pd.read_excel(arquivo_carregado).dropna(how="all")

    # --- CORREÇÃO AUTOMÁTICA ---
    df["nome"] = df["nome"].astype(str).str.strip().str.lower()

    if len(df.columns) > 3:
        coluna_d = df.columns[3]
        df["ritmo"] = df["ritmo"].fillna(df[coluna_d])

    df["tempo_limpo"] = df["tempo"].apply(limpar)
    df["ritmo_limpo"] = df["ritmo"].apply(limpar)

    df_agrupado = (
        df.groupby("nome")[["tempo_limpo", "ritmo_limpo"]]
        .mean()
        .reset_index()
    )
    df_agrupado["nome"] = df_agrupado["nome"].str.capitalize()

    df_agrupado["tempo_normal"] = (
        df_agrupado["tempo_limpo"] * df_agrupado["ritmo_limpo"]
    )
    df_agrupado["tempo_padrao"] = df_agrupado["tempo_normal"] * (
        1 + tolerancia
    )

    df_agrupado["pecas_hora"] = 3600 / df_agrupado["tempo_padrao"]
    df_agrupado["pecas_dia_teorica"] = df_agrupado["pecas_hora"] * jornada
    df_agrupado["pecas_dia_real"] = (
        df_agrupado["pecas_dia_teorica"] * oee_estimado
    )

    colunas_exibicao = [
        "nome",
        "tempo_limpo",
        "ritmo_limpo",
        "tempo_normal",
        "tempo_padrao",
        "pecas_hora",
        "pecas_dia_teorica",
        "pecas_dia_real",
    ]
    nomes_colunas = [
        "tarefa",
        "média tc (s)",
        "média ritmo",
        "tempo normal (s)",
        "tempo padrão (s)",
        "peças / hora",
        "cap. teórica / dia",
        "cap. real / dia (OEE)",
    ]

    if calcular_op:
        jornada_minutos = jornada * 60
        df_agrupado["tempo_padrao_min"] = df_agrupado["tempo_padrao"] / 60.0
        df_agrupado["operadores_nec"] = (
            df_agrupado["tempo_padrao_min"] * meta_producao
        ) / jornada_minutos
        colunas_exibicao.append("operadores_nec")
        nomes_colunas.append("nº operadores (necessários)")

    if calcular_maq:
        df_agrupado["maquinas_nec"] = (
            meta_producao / df_agrupado["pecas_dia_real"]
        )
        colunas_exibicao.append("maquinas_nec")
        nomes_colunas.append("nº máquinas (necessárias)")

    df_output = df_agrupado[colunas_exibicao].copy()
    df_output.columns = nomes_colunas

    # --- VISUALIZAÇÃO FORMATADA NA TELA DO NAVEGADOR ---
    df_visual = df_output.copy()
    formatos = {
        "média tc (s)": "{:.2f}",
        "média ritmo": "{:.2f}",
        "tempo normal (s)": "{:.2f}",
        "tempo padrão (s)": "{:.2f}",
        "peças / hora": "{:.1f}",
        "cap. teórica / dia": "{:.0f}",
        "cap. real / dia (OEE)": "{:.0f}",
        "nº operadores (necessários)": "{:.2f}",
        "nº máquinas (necessárias)": "{:.2f}",
    }
    for col, formato in formatos.items():
        if col in df_visual.columns:
            df_visual[col] = df_visual[col].map(formato.format)

    st.subheader("📋 Relatório Processado")
    st.dataframe(df_visual, use_container_width=True)

    total_op_arred = math.ceil(df_agrupado["operadores_nec"].sum()) if calcular_op else 0
    total_maq_arred = math.ceil(df_agrupado["maquinas_nec"].sum()) if calcular_maq else 0
    soma_tp_total = df_agrupado["tempo_padrao"].sum()

    st.markdown("### 🧮 Recursos Totais Necessários (Arredondados)")
    col1, col2, col3, col4 = st.columns(4)
    if calcular_op:
        col1.metric("Total de Operadores", f"{total_op_arred} pessoas")
    if calcular_maq:
        col2.metric("Total de Máquinas", f"{total_maq_arred} máquinas")
    col3.metric("Tempo Padrão Total", f"{soma_tp_total:.2f} s")
    col4.metric("Projeto / Setor", nome_projeto)

    # --- 📊 GERAÇÃO DO EXCEL COM ALTO CONTRASTE E BORDAS FORTES ---
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        workbook = writer.book
        worksheet = workbook.create_sheet(title="Dashboard PCP", index=0)
        worksheet.views.sheetView[0].showGridLines = True

        cor_cabecalho_fundo = "0F172A"
        cor_kpi_fundo = "F1F5F9"
        cor_kpi_texto = "1E40AF"
        cor_destaque_fundo = "EFF6FF"
        
        borda_preta_media = Side(style="medium", color="000000")
        borda_preta_fina = Side(style="thin", color="000000")
        borda_dupla_abaixo = Side(style="double", color="000000")
        borda_cinza_fina = Side(style="thin", color="CBD5E1")

        font_titulo = Font(name="Segoe UI", size=18, bold=True, color="0F172A")
        font_sub = Font(name="Segoe UI", size=10, italic=True, color="475569")
        
        font_kpi_label = Font(name="Segoe UI", size=9, bold=True, color="334155")
        font_kpi_val = Font(name="Segoe UI", size=18, bold=True, color=cor_kpi_texto)
        
        font_cabecalho = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        font_dados = Font(name="Segoe UI", size=10, color="000000", bold=False)
        font_dados_bold = Font(name="Segoe UI", size=10, color="000000", bold=True)

        fill_cabecalho = PatternFill(start_color=cor_cabecalho_fundo, end_color=cor_cabecalho_fundo, fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_destaque = PatternFill(start_color=cor_destaque_fundo, end_color=cor_destaque_fundo, fill_type="solid")
        fill_kpi = PatternFill(start_color=cor_kpi_fundo, end_color=cor_kpi_fundo, fill_type="solid")

        worksheet["B2"] = "PROGRAMAÇÃO E DIMENSIONAMENTO DE CAPACIDADE INDÚSTRIAL"
        worksheet["B2"].font = font_titulo
        worksheet["B3"] = f"Parâmetros aplicados | Tolerância: {tolerancia_pct}% | Turno: {jornada}h | OEE: {oee_pct}% | Meta: {meta_producao} peças"
        worksheet["B3"].font = font_sub

        # Bloco KPI 1: Operadores (Colunas B e C)
        if calcular_op:
            worksheet.merge_cells("B5:C5")
            worksheet.merge_cells("B6:C7")
            worksheet["B5"] = "TOTAL DE OPERADORES"
            worksheet["B6"] = f"{total_op_arred} Pessoas"
            
            for r in range(5, 8):
                for c in range(2, 4):
                    cell = worksheet.cell(row=r, column=c)
                    cell.fill = fill_kpi
                    t = borda_preta_media if r == 5 else None
                    b = borda_preta_media if r == 7 else None
                    l = borda_preta_media if c == 2 else None
                    rg = borda_preta_media if c == 3 else None
                    cell.border = Border(top=t, bottom=b, left=l, right=rg)
            
            worksheet["B5"].font = font_kpi_label
            worksheet["B5"].alignment = Alignment(horizontal="center", vertical="center")
            worksheet["B6"].font = font_kpi_val
            worksheet["B6"].alignment = Alignment(horizontal="center", vertical="center")

        # Bloco KPI 2: Máquinas (Colunas D e E)
        if calcular_maq:
            worksheet.merge_cells("D5:E5")
            worksheet.merge_cells("D6:E7")
            worksheet["D5"] = "TOTAL DE MÁQUINAS"
            worksheet["D6"] = f"{total_maq_arred} Conjuntos"
            
            for r in range(5, 8):
                for c in range(4, 6):
                    cell = worksheet.cell(row=r, column=c)
                    cell.fill = fill_kpi
                    t = borda_preta_media if r == 5 else None
                    b = borda_preta_media if r == 7 else None
                    l = borda_preta_media if c == 4 else None
                    rg = borda_preta_media if c == 5 else None
                    cell.border = Border(top=t, bottom=b, left=l, right=rg)

            worksheet["D5"].font = font_kpi_label
            worksheet["D5"].alignment = Alignment(horizontal="center", vertical="center")
            worksheet["D6"].font = font_kpi_val
            worksheet["D6"].alignment = Alignment(horizontal="center", vertical="center")

        # Bloco KPI 3: Tempo Padrão Total (Colunas F e G)
        worksheet.merge_cells("F5:G5")
        worksheet.merge_cells("F6:G7")
        worksheet["F5"] = "TEMPO PADRÃO TOTAL"
        worksheet["F6"] = f"{soma_tp_total:.2f} s"
        
        for r in range(5, 8):
            for c in range(6, 8):
                cell = worksheet.cell(row=r, column=c)
                cell.fill = fill_kpi
                t = borda_preta_media if r == 5 else None
                b = borda_preta_media if r == 7 else None
                l = borda_preta_media if c == 6 else None
                rg = borda_preta_media if c == 7 else None
                cell.border = Border(top=t, bottom=b, left=l, right=rg)

        worksheet["F5"].font = font_kpi_label
        worksheet["F5"].alignment = Alignment(horizontal="center", vertical="center")
        worksheet["F6"].font = font_kpi_val
        worksheet["F6"].alignment = Alignment(horizontal="center", vertical="center")

        # Bloco KPI 4: NOVO CARD - Nome do Projeto / Responsável (Colunas H e I)
        worksheet.merge_cells("H5:I5")
        worksheet.merge_cells("H6:I7")
        worksheet["H5"] = "PROJETO / SETOR"
        worksheet["H6"] = str(nome_projeto)
        
        for r in range(5, 8):
            for c in range(8, 10):
                cell = worksheet.cell(row=r, column=c)
                cell.fill = fill_kpi
                t = borda_preta_media if r == 5 else None
                b = borda_preta_media if r == 7 else None
                l = borda_preta_media if c == 8 else None
                rg = borda_preta_media if c == 9 else None
                cell.border = Border(top=t, bottom=b, left=l, right=rg)

        worksheet["H5"].font = font_kpi_label
        worksheet["H5"].alignment = Alignment(horizontal="center", vertical="center")
        worksheet["H6"].font = Font(name="Segoe UI", size=14, bold=True, color=cor_kpi_texto)
        worksheet["H6"].alignment = Alignment(horizontal="center", vertical="center")

        linha_inicio_tab = 10
        col_inicio_tab = 2

        for col_idx, col_name in enumerate(df_output.columns):
            cell = worksheet.cell(row=linha_inicio_tab, column=col_inicio_tab + col_idx)
            cell.value = col_name
            cell.font = font_cabecalho
            cell.fill = fill_cabecalho
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(top=borda_preta_media, bottom=borda_preta_media, left=borda_preta_fina, right=borda_preta_fina)
        worksheet.row_dimensions[linha_inicio_tab].height = 30

        total_linhas = len(df_output)
        for r_idx, row_data in enumerate(df_output.values):
            atual_row = linha_inicio_tab + 1 + r_idx
            worksheet.row_dimensions[atual_row].height = 22
            aplicar_zebra = r_idx % 2 == 1
            e_ultima_linha = r_idx == (total_linhas - 1)

            for c_idx, val in enumerate(row_data):
                cell = worksheet.cell(row=atual_row, column=col_inicio_tab + c_idx)
                cell.value = val
                cell.font = font_dados
                
                nome_coluna = df_output.columns[c_idx]
                if "necessários" in nome_coluna or "necessárias" in nome_coluna or "real" in nome_coluna:
                    cell.fill = fill_destaque
                    cell.font = font_dados_bold
                elif aplicar_zebra:
                    cell.fill = fill_zebra

                b_bottom = borda_dupla_abaixo if e_ultima_linha else borda_cinza_fina
                cell.border = Border(left=borda_cinza_fina, right=borda_cinza_fina, top=borda_cinza_fina, bottom=b_bottom)

                if c_idx == 0:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    
                    if nome_coluna in ["média tc (s)", "média ritmo", "tempo normal (s)", "tempo padrão (s)", "nº operadores (necessários)", "nº máquinas (necessárias)"]:
                        cell.number_format = "0.00"
                    elif nome_coluna == "peças / hora":
                        cell.number_format = "0.0"
                    elif nome_coluna in ["cap. teórica / dia", "cap. real / dia (OEE)"]:
                        cell.number_format = "#,##0"

        worksheet.column_dimensions["A"].width = 3
        for col_idx in range(len(df_output.columns)):
            col_letter = get_column_letter(col_inicio_tab + col_idx)
            max_len = len(df_output.columns[col_idx])
            for row in range(linha_inicio_tab, linha_inicio_tab + len(df_output) + 1):
                val_cell = worksheet.cell(row=row, column=col_inicio_tab + col_idx).value
                if val_cell:
                    max_len = max(max_len, len(str(val_cell)))
            worksheet.column_dimensions[col_letter].width = max(max_len + 5, 16)

    st.download_button(
        label="📥 Baixar Relatório Executivo de Alto Impacto",
        data=buffer.getvalue(),
        file_name="Painel_Cronoanalise_Contraste.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )